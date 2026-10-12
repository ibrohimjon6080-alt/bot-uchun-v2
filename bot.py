import asyncio
import os
from datetime import datetime
from aiogram import Bot, Dispatcher
from aiogram.filters import CommandStart, Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import Message, ReplyKeyboardMarkup, KeyboardButton, ReplyKeyboardRemove, FSInputFile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BOT_TOKEN = "8621333255:AAEVI2RyRMfSr70fKGoEyhVFMNhENp4u-h0"
EXCEL_FILE = "Sotuvchilar_Hisoboti.xlsx"

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()

class Form(StatesGroup):
    s1, s2, s3, s4, s5, s6, s7, s8, s9, s10, s11, s12, s13 = (State() for _ in range(13))

# Excel faylni yaratish va chiroyli dizayn berish
def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Hisobot"
        headers = [
            "Sana", "Klient nomi", "Mahsulot turi", "Brend nomi", 
            "Sifati haqida otziv", "Kg", "Narxi", "Yillik bonus", 
            "1 oyda o'rtacha abyom", "Qancha ast bilan ishlaydi", 
            "Tovarga garantiya", "Firma qulayliklari", "Nima uchun olmaysiz", 
            "2027-yildan ishlaydimi?"
        ]
        ws.append(headers)
        
        # Sarlavhalarga chiroyli dizayn (Ko'k orqa fon, oq qalin yozuv)
        fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        font = Font(bold=True, color="FFFFFF")
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        for col, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = fill
            cell.font = font
            cell.alignment = align
            # Ustunlar kengligini moslashtirish
            ws.column_dimensions[get_column_letter(col)].width = 22
            
        wb.save(EXCEL_FILE)

@dp.message(CommandStart())
async def start(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("Assalomu alaykum! Anketani to'ldirishni boshlaymiz.\n\n1. Klient nomini kiriting:", reply_markup=ReplyKeyboardRemove())
    await state.set_state(Form.s1)

# Xohlagan payt Excelni yuklab olish komandasi
@dp.message(Command("hisobot"))
async def send_excel(message: Message):
    if os.path.exists(EXCEL_FILE):
        file = FSInputFile(EXCEL_FILE)
        await message.answer_document(file, caption="📊 Barcha yig'ilgan ma'lumotlar Excel formatida!")
    else:
        await message.answer("Hozircha hech qanday ma'lumot yig'ilmagan.")

@dp.message(Form.s1)
async def process_s1(message: Message, state: FSMContext):
    await state.update_data(s1=message.text)
    kb = ReplyKeyboardMarkup(keyboard=[
        [KeyboardButton(text="Jitkiy gvoz"), KeyboardButton(text="Akrilik germetik")],
        [KeyboardButton(text="Tomchi"), KeyboardButton(text="Qopqoqlik pena")],
        [KeyboardButton(text="Professional Pena")]
    ], resize_keyboard=True, one_time_keyboard=True)
    await message.answer("2. Mahsulot turini tanlang:", reply_markup=kb)
    await state.set_state(Form.s2)

@dp.message(Form.s2)
async def process_s2(message: Message, state: FSMContext):
    await state.update_data(s2=message.text)
    await message.answer("3. Brend nomini kiriting:", reply_markup=ReplyKeyboardRemove())
    await state.set_state(Form.s3)

@dp.message(Form.s3)
async def process_s3(message: Message, state: FSMContext):
    await state.update_data(s3=message.text)
    await message.answer("4. Sifati haqida fikri (otziv):")
    await state.set_state(Form.s4)

@dp.message(Form.s4)
async def process_s4(message: Message, state: FSMContext):
    await state.update_data(s4=message.text)
    await message.answer("5. Kg (og'irligi):")
    await state.set_state(Form.s5)

@dp.message(Form.s5)
async def process_s5(message: Message, state: FSMContext):
    await state.update_data(s5=message.text)
    await message.answer("6. Narxi:")
    await state.set_state(Form.s6)

@dp.message(Form.s6)
async def process_s6(message: Message, state: FSMContext):
    await state.update_data(s6=message.text)
    await message.answer("7. Yillik bonus:")
    await state.set_state(Form.s7)

@dp.message(Form.s7)
async def process_s7(message: Message, state: FSMContext):
    await state.update_data(s7=message.text)
    await message.answer("8. 1 oyda o'rtacha abyom:")
    await state.set_state(Form.s8)

@dp.message(Form.s8)
async def process_s8(message: Message, state: FSMContext):
    await state.update_data(s8=message.text)
    await message.answer("9. Qancha ast (otstoy/qoldiq) bilan ishlaydi:")
    await state.set_state(Form.s9)

@dp.message(Form.s9)
async def process_s9(message: Message, state: FSMContext):
    await state.update_data(s9=message.text)
    await message.answer("10. Tovarga garantiya bormi / qanday?")
    await state.set_state(Form.s10)

@dp.message(Form.s10)
async def process_s10(message: Message, state: FSMContext):
    await state.update_data(s10=message.text)
    await message.answer("11. Firma qulayliklari:")
    await state.set_state(Form.s11)

@dp.message(Form.s11)
async def process_s11(message: Message, state: FSMContext):
    await state.update_data(s11=message.text)
    await message.answer("12. Bizdan nima uchun olmaysiz?")
    await state.set_state(Form.s12)

@dp.message(Form.s12)
async def process_s12(message: Message, state: FSMContext):
    await state.update_data(s12=message.text)
    await message.answer("13. Shu muammolarni bartaraf qilsak 2027 yildan biz bilan ishlaysizmi?")
    await state.set_state(Form.s13)

@dp.message(Form.s13)
async def process_s13(message: Message, state: FSMContext):
    await state.update_data(s13=message.text)
    data = await state.get_data()
    
    init_excel()
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active
    sana = datetime.now().strftime("%d.%m.%Y %H:%M")
    
    row = [
        sana, data['s1'], data['s2'], data['s3'], data['s4'], data['s5'], 
        data['s6'], data['s7'], data['s8'], data['s9'], data['s10'], 
        data['s11'], data['s12'], data['s13']
    ]
    ws.append(row)
    
    # Qatorlarni o'rtaga to'g'irlash
    align = Alignment(horizontal="center", vertical="center")
    for cell in ws[ws.max_row]:
        cell.alignment = align
        
    wb.save(EXCEL_FILE)
    
    await message.answer("✅ Rahmat! Ma'lumotlar saqlandi.\n\nTayyor Excel faylni olish uchun /hisobot deb yozing.")
    await state.clear()

async def main():
    init_excel()
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
